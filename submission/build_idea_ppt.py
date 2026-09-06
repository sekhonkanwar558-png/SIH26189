"""Build the SIH 2026 idea presentation for PS 26189 from the official template.

The template is the one on sih.gov.in and the one Dr. Kanu Goel attached to the
(since-deleted) Classroom post. Her three instructions are honoured here:

  * the last slide (IMPORTANT INSTRUCTIONS) is deleted
  * Team ID is left blank -- the portal assigns it after SPOC nomination
  * Team Name is whatever was typed on the internal-hackathon registration form

The template's own rules are honoured too: six slides including the title, the
slide titles and the idea-detail pointers kept verbatim, and -- the reason
slides 2 and 3 are two-column -- "post your idea in points / diagrams /
infographics / pictures". Slide 3's pointer asks for flow charts, images or the
working prototype by name, so it gets a drawn flow chart; slide 2 describes the
prototype, so it carries a screenshot of it running.

    python submission/build_idea_ppt.py

Output: submission/<TEAM_NAME>.pptx
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# --- must match the internal-hackathon registration form exactly --------------
TEAM_NAME = "Suishōdama"

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "SIH2026-IDEA-Presentation-Format.pptx"
SCREENSHOT = HERE / "assets" / "prototype-chat-and-brain.png"
OUTPUT = HERE / f"{TEAM_NAME}.pptx"

INK = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT = RGBColor(0x1F, 0x38, 0x64)
MUTED = RGBColor(0x3D, 0x3D, 0x3D)
BOX_FILL = RGBColor(0xEE, 0xF2, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# The usable band between the title zone and the footer bar.
TOP, BOTTOM = 1.36, 6.88

TITLE_FIELDS = [
    ("Problem Statement ID – ", "26189"),
    ("Problem Statement Title- ", "AI-Powered Criminal Network Analysis System"),
    ("Theme- ", "Blockchain & Cybersecurity"),
    ("PS Category- ", "Software"),
    ("Team ID- ", ""),  # blank on purpose: assigned by the portal
    ("Team Name (Registered on portal)- ", TEAM_NAME),
]

# (level, kind, text) -- kind is lead | pointer | sub | bullet | note | caption
SLIDE2_LEAD = [
    (0, "lead", "Suishōdama — a per-case investigative agent that reasons over the case's own "
                "knowledge graph, not a context window."),
]
SLIDE2_LEFT = [
    (0, "pointer", "Proposed Solution (Describe your Idea/Solution/Prototype)"),
    (0, "sub", "Detailed explanation of the proposed solution"),
    (1, "bullet", "An officer opens a case, and that case gets an agent of its own — its own knowledge "
                  "graph, analytics, memory and hash-chained chain of custody, in its own file on disk."),
    (1, "bullet", "He drops in FIRs, CDRs, bank statements, surveillance logs, social exports, criminal "
                  "history and intelligence reports — extracted deterministically, with no model anywhere "
                  "in the ingest path."),
    (1, "bullet", "He asks in English. The answer arrives with the route through the graph lit hop by "
                  "hop, and every claim traceable to the character offset in the document it was read "
                  "out of."),
    (0, "sub", "How it addresses the problem"),
    (1, "bullet", "NCRB's stated pain is that the data is already on the desk and the links are still "
                  "missed. Seven source kinds are named in the statement, and there is a handler for "
                  "all seven."),
    (1, "bullet", "“Key influencers” is computed, not guessed — betweenness on a person-projected "
                  "graph, Louvain communities, and hidden brokers ranked person-against-person."),
]
SLIDE2_RIGHT = [
    (0, "sub", "Innovation and uniqueness of the solution"),
    (1, "bullet", "Not a dashboard — a visualisation cannot be asked anything."),
    (1, "bullet", "Not a chat with documents in it — the prompt is O(1) in the size of the case."),
    (1, "bullet", "Isolation by construction: a different case is a different file."),
]
SLIDE2_CAPTION = ("The working prototype — one chat, the case's brain beside it, and the cited source "
                  "open at its character offsets.")

SLIDE3_LEFT = [
    (0, "pointer", "Technologies to be used (e.g. programming languages, frameworks, hardware)"),
    (1, "bullet", "Backend — Python 3.11+, FastAPI, SQLite (one file per case), NetworkX, Anthropic API."),
    (1, "bullet", "Extraction — readers for PDF / CSV / prose, deterministic identifier patterns, "
                  "cue-based NER, optional spaCy."),
    (1, "bullet", "Analytics — degree and betweenness centrality, Louvain communities, anomaly and "
                  "finding detection; seeded and cached."),
    (1, "bullet", "Interface — React 19, TypeScript, Vite, Tailwind v4, Cytoscape.js, TanStack Query."),
    (1, "bullet", "Integrity — SHA-256 hash-chained, append-only chain of custody, one per case."),
    (1, "bullet", "No Neo4j and no cloud graph service: per-case isolation becomes a file rather than "
                  "a permissions model."),
    (0, "note", "Working prototype today: 22 case-scoped API endpoints, 38 tests passing, and a demo "
                "case of 125 nodes and 1,148 links built from all seven source kinds."),
]
SLIDE3_POINTER = "Methodology and process for implementation (Flow Charts/Images/ working prototype)"
FLOW = [
    ("DOCUMENTS", "FIRs · CDRs · financial · surveillance · social · history · intel"),
    ("1.  INGEST", "readers → patterns → NER → pipeline  ·  no model in this path"),
    ("2.  GRAPH", "one SQLite file per case  ·  stable ids  ·  idempotent merge"),
    ("3.  ANALYTICS", "centrality · Louvain · anomalies · findings the model never invents"),
    ("4.  AGENT", "13 tools  ·  decides what to ask the brain, holds none of it"),
    ("5.  ANSWER", "every citation verified in code  ·  the path lit in the graph"),
]
FLOW_BAND = "CHAIN OF CUSTODY — hash-linked, append-only, one per case"

SLIDES_FULL = {
    4: [
        (0, "pointer", "Analysis of the feasibility of the idea"),
        (1, "bullet", "Built and running rather than proposed — 38 tests pass, 22 endpoints, 125 nodes "
                      "and 1,148 links across all seven source kinds."),
        (1, "bullet", "Ingest is linear and measured — 106 KB of prose in 3.8 s, down from 63 s; a "
                      "50,000-row CDR in about 90 s."),
        (1, "bullet", "Runs on a single machine with no external service — SQLite and NetworkX, nothing "
                      "to install, authenticate or keep up."),
        (1, "bullet", "$0.0235 per question, down from $0.10, through prompt caching with the system "
                      "prompt split into three blocks."),
        (0, "pointer", "Potential challenges and risks"),
        (1, "bullet", "Police data cannot leave the premises."),
        (1, "bullet", "No key or no network means no assistant."),
        (1, "bullet", "A scanned document has no text layer to read."),
        (1, "bullet", "An estimate read out as a measurement would mislead an investigation."),
        (0, "pointer", "Strategies for overcoming these challenges"),
        (1, "bullet", "The model's whole surface is two functions — an on-premises model drops in "
                      "untouched, and everything but the assistant already runs with no model at all."),
        (1, "bullet", "Nothing degrades silently — no fallback, no offline impostor. A degraded system "
                      "the officer cannot tell from a working one is worse than an outage."),
        (1, "bullet", "A document that could not be read says so, instead of registering with a green "
                      "tick, a valid custody entry and zero entities."),
        (1, "bullet", "Betweenness is exact below 500 projected nodes and estimated above — and the "
                      "answer says which it used."),
    ],
    5: [
        (0, "pointer", "Potential impact on the target audience"),
        (1, "bullet", "The connection the officer would never have written down: the most central "
                      "person in the demo case appears in one document only, a call log, and 22% of "
                      "every shortest path runs through him."),
        (1, "bullet", "Asked to volunteer, it named a woman in no FIR and not in the ground-truth "
                      "file — four Ludhiana towers, three Delhi, in call contact with both principals. "
                      "A bridge nobody planted."),
        (1, "bullet", "The engine is crime-type agnostic: nothing in ingest, the graph, analytics or "
                      "the agent knows what trafficking is, so the same system serves NDPS, financial "
                      "fraud or organised crime."),
        (0, "pointer", "Benefits of the solution (social, economic, environmental, etc.)"),
        (1, "bullet", "Social — for NCRB's Women Safety Division, faster identification of traffickers "
                      "and repeat offenders, and networks read across a case no officer could hold in "
                      "his head."),
        (1, "bullet", "Evidentiary — every claim traceable to a character offset, and a hash-chained "
                      "custody log means history cannot be quietly rewritten. Break an entry and it "
                      "names the entry and the reason."),
        (1, "bullet", "Economic — one machine, no licence, no cluster, no graph-database contract, "
                      "and $0.0235 a question."),
        (1, "bullet", "Operational — one chat, learned in seconds by a non-technical user, and "
                      "knowledge that compounds over months instead of resetting with the thread."),
    ],
    6: [
        (0, "pointer", "Details / Links of the reference and research work"),
        (1, "bullet", "Problem statement 26189, Ministry of Home Affairs / NCRB (Women Safety "
                      "Division) — sih.gov.in/sih2026PS"),
        (1, "bullet", "U. Brandes, “A faster algorithm for betweenness centrality”, Journal of "
                      "Mathematical Sociology 25(2), 2001 — the centrality measure used to identify "
                      "key influencers."),
        (1, "bullet", "V. Blondel et al., “Fast unfolding of communities in large networks”, J. Stat. "
                      "Mech., 2008 — the Louvain method used for community detection. "
                      "arxiv.org/abs/0803.0476"),
        (1, "bullet", "NetworkX — graph algorithms and sampled-pivot betweenness. networkx.org"),
        (1, "bullet", "NATO Admiralty Code (source reliability × information credibility) — the "
                      "grading applied to intelligence reports, discounting what they imply."),
        (1, "bullet", "FIPS 180-4 (SHA-256) — the hash used for the append-only chain of custody."),
        (1, "bullet", "Anthropic — tool use and prompt caching. docs.anthropic.com"),
        (1, "bullet", "Cytoscape.js — graph rendering in the browser. js.cytoscape.org"),
    ],
}

#        size, bold, colour, space_before, indent
STYLE = {
    "lead":    (15, True, ACCENT, 0, 0.00),
    "pointer": (14.5, True, ACCENT, 8, 0.00),
    "sub":     (12.5, True, MUTED, 7, 0.10),
    "bullet":  (11, False, INK, 3, 0.26),
    "note":    (11.5, True, MUTED, 9, 0.06),
    "caption": (9.5, False, MUTED, 4, 0.00),
}


def delete_slide(prs, index):
    """Remove a slide by index -- python-pptx has no API for this."""
    slide_id = prs.slides._sldIdLst[index]
    prs.part.drop_rel(slide_id.rId)
    prs.slides._sldIdLst.remove(slide_id)


def shape_named(slide, name):
    for shape in slide.shapes:
        if shape.name == name:
            return shape
    return None


def style_paragraph(para, kind, level=0):
    """Apply a STYLE entry, and kill the template's inherited list glyphs.

    The template's list style puts a bullet on some levels and not others,
    which showed up as a stray dot on the first pointer of each slide. buNone
    turns them all off; the em-dashes carry the hierarchy instead.
    """
    size, bold, colour, space_before, indent = STYLE[kind]
    para.level = level
    # The template justifies body text, which stretches word spacing on any
    # line that wraps -- visible on every pointer. Left-align instead.
    para.alignment = PP_ALIGN.LEFT
    para.space_before = Pt(space_before)
    para.space_after = Pt(1)
    pPr = para._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent))))
    pPr.set("indent", "0")
    for tag in ("a:buChar", "a:buAutoNum", "a:buNone"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    pPr.append(pPr.makeelement(qn("a:buNone"), {}))
    return size, bold, colour


def fill_rows(tf, rows, size_delta=0.0):
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    paragraphs = tf.paragraphs
    for para in paragraphs[1:]:
        para._p.getparent().remove(para._p)
    first = paragraphs[0]
    for run in list(first.runs):
        run._r.getparent().remove(run._r)

    for i, (level, kind, text) in enumerate(rows):
        para = first if i == 0 else tf.add_paragraph()
        size, bold, colour = style_paragraph(para, kind, level)
        run = para.add_run()
        run.text = ("—  " + text) if kind == "bullet" else text
        run.font.name = "Arial"
        run.font.size = Pt(size + size_delta)
        run.font.bold = bold
        run.font.color.rgb = colour


def move(shape, left, top, width, height):
    shape.left, shape.top = Inches(left), Inches(top)
    shape.width, shape.height = Inches(width), Inches(height)


def draw_flowchart(slide, left, top, width, height):
    """The methodology, drawn -- which is what the template's pointer asks for."""
    band_h, gap = 0.38, 0.15
    n = len(FLOW)
    box_h = (height - band_h - gap - (n - 1) * gap) / n

    y = top
    for i, (title, subtitle) in enumerate(FLOW):
        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(y), Inches(width), Inches(box_h)
        )
        box.fill.solid()
        box.fill.fore_color.rgb = BOX_FILL
        box.line.color.rgb = ACCENT
        box.line.width = Pt(0.75)
        box.shadow.inherit = False

        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Inches(0.06)
        tf.margin_top = tf.margin_bottom = 0

        head = tf.paragraphs[0]
        head.alignment = PP_ALIGN.CENTER
        head.space_after = Pt(0)
        run = head.add_run()
        run.text = title
        run.font.name, run.font.size, run.font.bold = "Arial", Pt(10), True
        run.font.color.rgb = ACCENT

        sub = tf.add_paragraph()
        sub.alignment = PP_ALIGN.CENTER
        sub.space_before = Pt(1)
        run = sub.add_run()
        run.text = subtitle
        run.font.name, run.font.size = "Arial", Pt(7.5)
        run.font.color.rgb = MUTED

        if i < n - 1:
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.ISOSCELES_TRIANGLE,
                Inches(left + width / 2 - 0.09), Inches(y + box_h + 0.02),
                Inches(0.18), Inches(gap - 0.04),
            )
            arrow.rotation = 180
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = ACCENT
            arrow.line.fill.background()
            arrow.shadow.inherit = False
        y += box_h + gap

    band = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(y), Inches(width), Inches(band_h)
    )
    band.fill.solid()
    band.fill.fore_color.rgb = ACCENT
    band.line.fill.background()
    band.shadow.inherit = False
    tf = band.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_top = tf.margin_bottom = 0
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    run = para.add_run()
    run.text = FLOW_BAND
    run.font.name, run.font.size, run.font.bold = "Arial", Pt(8.5), True
    run.font.color.rgb = WHITE


def build_slide2(slide):
    lead = shape_named(slide, "TextBox 8")
    move(lead, 0.55, TOP - 0.08, 12.25, 0.58)
    fill_rows(lead.text_frame, SLIDE2_LEAD)

    left_box = slide.shapes.add_textbox(
        Inches(0.55), Inches(1.96), Inches(6.45), Inches(BOTTOM - 1.96)
    )
    fill_rows(left_box.text_frame, SLIDE2_LEFT)

    pic_left, pic_width = 7.35, 5.95
    pic_height = pic_width / 2.123
    slide.shapes.add_picture(
        str(SCREENSHOT), Inches(pic_left), Inches(1.96),
        width=Inches(pic_width), height=Inches(pic_height),
    )
    cap_top = 1.96 + pic_height + 0.04
    cap = slide.shapes.add_textbox(
        Inches(pic_left), Inches(cap_top), Inches(pic_width), Inches(0.52)
    )
    fill_rows(cap.text_frame, [(0, "caption", SLIDE2_CAPTION)])

    right_top = cap_top + 0.56
    right = slide.shapes.add_textbox(
        Inches(pic_left), Inches(right_top), Inches(pic_width), Inches(BOTTOM - right_top)
    )
    fill_rows(right.text_frame, SLIDE2_RIGHT, size_delta=-0.5)


def build_slide3(slide):
    left_box = shape_named(slide, "TextBox 8")
    move(left_box, 0.55, TOP, 6.60, BOTTOM - TOP)
    fill_rows(left_box.text_frame, SLIDE3_LEFT)

    head = slide.shapes.add_textbox(
        Inches(7.45), Inches(TOP - 0.04), Inches(5.50), Inches(0.62)
    )
    fill_rows(head.text_frame, [(0, "pointer", SLIDE3_POINTER)], size_delta=-2.5)

    draw_flowchart(slide, 7.45, TOP + 0.62, 5.50, BOTTOM - TOP - 0.62)


def build_title_page(slide):
    box = shape_named(slide, "TextBox 9")
    move(box, 0.36, 2.05, 6.10, 5.20)
    tf = box.text_frame
    tf.word_wrap = True
    paragraphs = tf.paragraphs
    for para in paragraphs[1:]:
        para._p.getparent().remove(para._p)
    for run in list(paragraphs[0].runs):
        run._r.getparent().remove(run._r)

    for i, (label, value) in enumerate(TITLE_FIELDS):
        para = paragraphs[0] if i == 0 else tf.add_paragraph()
        para.space_after = Pt(8)
        for text, bold in ((label, True), (value, False)):
            if not text:
                continue
            run = para.add_run()
            run.text = text
            run.font.name = "Arial"
            run.font.size = Pt(17)
            run.font.bold = bold


def remove_footer(slide):
    """Drop the template's "@SIH Idea submission- Template" footer.

    His call: the blue bar keeps its slide number and nothing else. The whole
    placeholder goes rather than just its text -- an emptied placeholder is
    still a placeholder, and PowerPoint can repopulate one from the layout.
    """
    footer = shape_named(slide, "Footer Placeholder 6")
    if footer is not None:
        footer._element.getparent().remove(footer._element)


def stamp_team_name(slide):
    for shape in slide.shapes:
        if shape.name.startswith("Oval") and shape.has_text_frame:
            tf = shape.text_frame
            tf.word_wrap = True
            para = tf.paragraphs[0]
            for run in list(para.runs):
                run._r.getparent().remove(run._r)
            run = para.add_run()
            run.text = TEAM_NAME
            run.font.name, run.font.size, run.font.bold = "Arial", Pt(9), True


def main():
    prs = Presentation(TEMPLATE)
    delete_slide(prs, 6)  # her instruction, and the template's own

    build_title_page(prs.slides[0])
    build_slide2(prs.slides[1])
    build_slide3(prs.slides[2])
    for number, rows in SLIDES_FULL.items():
        slide = prs.slides[number - 1]
        box = shape_named(slide, "TextBox 8")
        move(box, 0.55, TOP, 12.25, BOTTOM - TOP)
        fill_rows(box.text_frame, rows)

    for number in range(2, 7):
        stamp_team_name(prs.slides[number - 1])
        remove_footer(prs.slides[number - 1])

    prs.save(OUTPUT)
    print(f"wrote {OUTPUT}  ({len(prs.slides._sldIdLst)} slides)")


if __name__ == "__main__":
    main()
